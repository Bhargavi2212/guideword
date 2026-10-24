"""The review and export routes on a live $0 fake run (spec 3.3, 3.5; Phase 3 plan PR 6, G2, G5,
Q17). Every model call goes through the fake gateway; the run reaches LOPA, chains and
recommendations."""

import io
import json
import re
from collections.abc import Callable
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from urllib.parse import quote

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session, sessionmaker

from guideword.api import models as api
from guideword.api.app import NOT_BUILT_PHASE, create_app
from guideword.api.mock import WORKSHEET_COLUMNS
from guideword.config import load_config
from guideword.gateway.client import Gateway
from guideword.gateway.fake import fake_gateway, reply
from guideword.gateway.transport import Reply, ToolCallRequest
from guideword.runs import db
from guideword.schema.enums import BudgetUse
from guideword.settings import Settings

REPO = Path(__file__).resolve().parents[3]
EXAMPLE = REPO / "backend" / "tests" / "schema" / "fixtures" / "example_mixing_unit.json"
NOW = datetime(2026, 10, 20, 9, 0, tzinfo=UTC)
SECRET = "test-advance-secret"
GROUPS = 6
BUILT_IN_PR_6 = {
    ("GET", "/runs/{id}/shortlist"),
    ("GET", "/runs/{id}/worksheet"),
    ("POST", "/runs/{id}/overrides"),
    ("GET", "/runs/{id}/export.xlsx"),
    ("GET", "/runs/{id}/export/coverage.json"),
    ("GET", "/runs/{id}/export/run.json"),
}
_BLOCK = r"<{0}>\s*(.*?)\s*</{0}>"


def _block(request: dict[str, Any], name: str) -> Any:
    match = re.search(_BLOCK.format(name), request["messages"][0]["content"], re.S)
    assert match is not None, name
    return json.loads(match.group(1))


def _group(question_id: str) -> int:
    return sum(map(ord, question_id)) % GROUPS


def script(job: str, request: dict[str, Any]) -> Reply:
    """One outside-library finding per question on the first location, its cause by group;
    severity S4 for even groups, S3 for odd; the connector chains the first two lines; every
    layer applies; one recommendation per line; the last group is not credible."""
    if job == "section_suggestion":
        return reply({"suggestions": []})
    if job == "plant_screen":
        return reply({"hypotheses": []})
    if job == "orchestrator":
        view = _block(request, "view")
        agents = sorted(view["open"])
        if agents:
            retry = view["retry_at_hard"].get(agents[0], {"count": 0})["count"]
            level = "bulk" if view["open"][agents[0]]["count"] > retry else "hard"
            name, args = "dispatch_specialist", {"agent": agents[0], "level": level}
        else:
            name, args = "finish", {}
        return reply(tool_calls=(ToolCallRequest("t", name, json.dumps({"reason": "go", **args})),))
    if job == "severity_reason":
        rated = [
            {
                "finding_id": f["finding_id"],
                "severity_class": "S4" if int(f["cause"].split()[-1]) % 2 == 0 else "S3",
                "reason": "test",
            }
            for f in _block(request, "findings")
        ]
        return reply({"severities": rated})
    if job == "connector":
        lines = _block(request, "scenarios")
        links = [lines[0]["handle"], lines[1]["handle"]]
        return reply({"chains": [{"links": links, "description": "one then two", "reason": "r"}]})
    if job == "chain_severity":
        chains = _block(request, "chains")
        return reply(
            {
                "severities": [
                    {"chain": c["handle"], "severity_class": "S4", "reason": "r"} for c in chains
                ]
            }
        )
    if job == "lopa_argument":
        lines = [
            {
                "target": line["handle"],
                "layers": [
                    {"layer_id": layer["layer_id"], "applies": True, "reason": "it applies"}
                    for layer in line["layers"]
                ],
            }
            for line in _block(request, "lines")
        ]
        return reply({"lines": lines})
    if job == "recommendation":
        recs = [
            {"target": line["handle"], "action": "add an independent trip", "reason": "r"}
            for line in _block(request, "lines")
        ]
        return reply({"recommendations": recs})
    locations = _block(request, "context").get("locations") or []
    questions = _block(request, "questions")
    if not locations:
        answers = [
            {"question_id": q["id"], "not_credible": {"reason": "no location"}, "findings": []}
            for q in questions
        ]
    else:
        # the last group is dismissed as not credible, so the worksheet has such rows
        answers = [
            {"question_id": q["id"], "not_credible": {"reason": "no path"}, "findings": []}
            if _group(q["id"]) == GROUPS - 1
            else {
                "question_id": q["id"],
                "not_credible": None,
                "findings": [
                    {
                        "cause_id": "outside_library",
                        "cause_text": f"outside {_group(q['id'])}",
                        "failing_component_type": "external",
                        "failing_component_ids": [locations[0]["id"]],
                        "consequence": "a release",
                        "safeguard_ids": [],
                        "adequacy": "unknown",
                        "confidence": "sure",
                        "reason": "test",
                    }
                ],
            }
            for q in questions
        ]
    return reply({"answers": answers, "added_questions": []})


class Api:
    def __init__(self, tmp_path: Path) -> None:
        self.database = f"sqlite:///{(tmp_path / 'g.db').as_posix()}"
        settings = Settings(
            _env_file=None,  # type: ignore[call-arg]
            guideword_mock=False,
            database_url=self.database,
            internal_advance_secret=SECRET,
        )
        config = load_config(global_budget_usd=25.0)

        def gateway(sessions: sessionmaker[Session], use: BudgetUse) -> Gateway:
            built, _ = fake_gateway(
                config,
                script,
                sessions=sessions,
                ledger_path=tmp_path / "l.jsonl",
                clock=lambda: NOW,
                use=use,
                phase="5",
            )
            return built

        factory: Callable[[sessionmaker[Session], BudgetUse], Gateway] = gateway
        self.client = TestClient(
            create_app(settings=settings, clock=lambda: NOW, gateway_factory=factory)
        )

    def finished_run(self) -> str:
        data = json.loads(EXAMPLE.read_text(encoding="utf-8"))
        plant = self.client.post("/plants", json=data).json()["plant_record_id"]
        url = f"/plants/{quote(plant, safe='')}"
        proposal = self.client.post(f"{url}/sections/propose").json()
        self.client.put(f"{url}/sections", json={"sections": proposal["sections"]})
        run_id = str(
            self.client.post(
                "/runs",
                json={
                    "plant_id": plant,
                    "spend_cap_usd": 0.2,
                    "mode": "full",
                    "use": "development",
                    "post_spend_cap_usd": 0.1,
                },
            ).json()["run_id"]
        )
        advance = f"/internal/runs/{quote(run_id, safe='')}/advance"
        assert self.client.post(advance, headers={"X-Internal-Secret": SECRET}).status_code == 202
        assert self.get(run_id, "").json()["run"]["status"] == "finished"
        return run_id

    def get(self, run_id: str, rest: str) -> Any:
        return self.client.get(f"/runs/{quote(run_id, safe='')}{rest}")

    def override(self, run_id: str, target: str, action: str, value: Any) -> Any:
        body = {"target_id": target, "action": action, "new_value": value, "reason": "r", "by": "B"}
        return self.client.post(f"/runs/{quote(run_id, safe='')}/overrides", json=body)

    def shortlist(self, run_id: str) -> api.Shortlist:
        response = self.get(run_id, "/shortlist")
        assert response.status_code == 200, response.text
        return api.Shortlist.model_validate(response.json())


@pytest.fixture
def app(tmp_path: Path) -> Api:
    return Api(tmp_path)


@pytest.fixture
def run_id(app: Api) -> str:
    return app.finished_run()


def test_the_built_routes_left_not_built_phase() -> None:
    assert not BUILT_IN_PR_6 & set(NOT_BUILT_PHASE)
    assert set(NOT_BUILT_PHASE) == {("GET", "/demo")}


def test_the_shortlist_serves_the_runs_lopa_chains_and_recommendations(
    app: Api, run_id: str
) -> None:
    items = app.shortlist(run_id).items
    assert items and [i.rank for i in items] == list(range(1, len(items) + 1))
    assert any(i.target_kind == "chain" and i.chain is not None for i in items)
    assert any(i.lopa is not None for i in items)
    assert any(i.recommendations for i in items)
    assert all(i.evidence == [] for i in items)


def test_change_credit_recomputes_gap_on_read(app: Api, run_id: str) -> None:
    item = next(
        i
        for i in app.shortlist(run_id).items
        if i.lopa and any(layer.credited for layer in i.lopa.layers)
    )
    assert item.lopa is not None and item.gap is not None
    layer = next(layer for layer in item.lopa.layers if layer.credited)
    assert layer.pfd is not None
    response = app.override(
        run_id, item.lopa.id, "change_credit", {"layer_id": layer.layer_id, "credited": False}
    )
    assert response.status_code == 200, response.text
    result = api.OverrideResult.model_validate(response.json())
    assert result.stored is True and result.override.replaces is None
    seen = next(i for i in app.shortlist(run_id).items if i.target_id == item.target_id)
    assert seen.gap == pytest.approx(item.gap / layer.pfd)
    # the stored record is unchanged: the run dump has it as made
    dump = api.RunDump.model_validate(app.get(run_id, "/export/run.json").json())
    stored = next(r for r in dump.lopa_records if r.id == item.lopa.id)
    assert stored.gap == pytest.approx(item.gap)
    assert [o.id for o in dump.overrides] == [result.override.id]


def test_change_severity_moves_tolerable_and_the_second_override_replaces_the_first(
    app: Api, run_id: str
) -> None:
    item = next(i for i in app.shortlist(run_id).items if i.target_kind == "finding" and i.lopa)
    first = app.override(run_id, item.target_id, "change_severity", {"severity_class": "S5"})
    second = app.override(run_id, item.target_id, "change_severity", {"severity_class": "S2"})
    assert first.status_code == second.status_code == 200
    assert second.json()["override"]["replaces"] == first.json()["override"]["id"]
    seen = next(i for i in app.shortlist(run_id).items if i.target_id == item.target_id)
    assert seen.finding is not None and seen.finding.severity_class == "S2"
    assert seen.lopa is not None and seen.lopa.severity_class == "S2"


def test_override_on_a_recorded_run_is_403(app: Api, run_id: str) -> None:
    sessions = db.make_session_factory(app.database)
    with sessions.begin() as session:
        row = session.scalars(select(db.RunRow).where(db.RunRow.id == run_id)).one()
        row.record = {**row.record, "label": {"kind": "recorded", "recorded_at": NOW.isoformat()}}
    target = app.shortlist(run_id).items[0].target_id
    assert app.override(run_id, target, "accept", None).status_code == 403


def test_edit_of_any_other_field_is_422(app: Api, run_id: str) -> None:
    item = next(i for i in app.shortlist(run_id).items if i.target_kind == "finding")
    assert app.override(run_id, item.target_id, "edit", {"severity_reason": "x"}).status_code == 422
    assert app.override(run_id, item.target_id, "edit", {"action": "x"}).status_code == 422
    assert app.override(run_id, item.target_id, "edit", {"consequence": "x"}).status_code == 200


def test_answer_on_a_non_engineer_question_is_422(app: Api, run_id: str) -> None:
    questions = app.get(run_id, "/questions").json()["questions"]
    matrix = next(q for q in questions if q["kind"] != "engineer")
    assert app.override(run_id, matrix["id"], "answer", {"text": "x"}).status_code == 422


def test_an_unknown_target_is_404_and_a_wrong_layer_422(app: Api, run_id: str) -> None:
    assert app.override(run_id, "f:nope", "accept", None).status_code == 404
    item = next(i for i in app.shortlist(run_id).items if i.lopa is not None)
    assert item.lopa is not None
    wrong = {"layer_id": "NOT-A-LAYER", "credited": True}
    assert app.override(run_id, item.lopa.id, "change_credit", wrong).status_code == 422


def test_answered_engineer_question_counts_as_checked_on_read(app: Api, run_id: str) -> None:
    section = app.get(run_id, "/sections").json()["sections"][0]["id"]
    added = app.client.post(
        f"/runs/{quote(run_id, safe='')}/questions",
        json={"section_id": section, "text": "What if the lid is left open?"},
    ).json()["question"]
    before = api.Coverage.model_validate(app.get(run_id, "/coverage").json())
    engineer = next(line for line in before.lists if line.list == "engineer")
    assert (
        app.override(run_id, added["id"], "answer", {"text": "vapour in the hall"}).status_code
        == 200
    )
    for path in ("/coverage", "/export/coverage.json"):
        after = api.Coverage.model_validate(app.get(run_id, path).json())
        line = next(line for line in after.lists if line.list == "engineer")
        assert line.checked == engineer.checked + 1
    rows = api.Worksheet.model_validate(app.get(run_id, "/worksheet").json()).rows
    row = next(r for r in rows if r.row_id == added["id"])
    assert row.status == "answered" and row.consequences == "vapour in the hall"


def test_the_record_route_knows_chains_lopa_recommendations_and_overrides(
    app: Api, run_id: str
) -> None:
    items = app.shortlist(run_id).items
    chain = next(i.chain for i in items if i.chain is not None)
    lopa = next(i.lopa for i in items if i.lopa is not None and i.lopa.model_call_id)
    rec = next(r for i in items for r in i.recommendations)
    made = app.override(run_id, rec.id, "accept", None).json()["override"]

    def detail(record_id: str) -> api.RecordDetail:
        response = app.get(run_id, f"/records/{quote(record_id, safe='')}")
        assert response.status_code == 200, response.text
        return api.RecordDetail.model_validate(response.json())

    found = detail(chain.id)
    assert found.record_kind == "chain" and found.parent_ids == chain.links
    assert [c.id for c in found.model_calls] == [chain.model_call_id]
    found = detail(lopa.id)
    assert found.record_kind == "lopa" and found.parent_ids == [lopa.target_id]
    assert [c.id for c in found.model_calls] == [lopa.model_call_id]
    found = detail(rec.id)
    assert found.record_kind == "recommendation" and found.parent_ids == [rec.target_id]
    assert [c.id for c in found.model_calls] == [rec.model_call_id]
    found = detail(made["id"])
    assert found.record_kind == "override" and found.parent_ids == [rec.id]


def test_xlsx_has_the_step_8_columns_one_row_per_worksheet_row_overrides_applied(
    app: Api, run_id: str
) -> None:
    item = next(i for i in app.shortlist(run_id).items if i.target_kind == "finding")
    edited = app.override(run_id, item.target_id, "edit", {"consequence": "the hall floods"})
    assert edited.status_code == 200
    rows = api.Worksheet.model_validate(app.get(run_id, "/worksheet").json()).rows
    response = app.get(run_id, "/export.xlsx")
    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert sheet is not None
    table = list(sheet.iter_rows(values_only=True))
    assert list(table[0]) == [title for title, _ in WORKSHEET_COLUMNS]
    assert len(table) - 1 == len(rows)
    column = [title for title, _ in WORKSHEET_COLUMNS].index("Consequences")
    assert "the hall floods" in [r[column] for r in table[1:]]


def test_exports_match_the_contract(app: Api, run_id: str) -> None:
    dump = api.RunDump.model_validate(app.get(run_id, "/export/run.json").json())
    assert dump.run.id == run_id and dump.chains and dump.lopa_records and dump.recommendations
    assert dump.label.kind == "live"
    api.Coverage.model_validate(app.get(run_id, "/export/coverage.json").json())
    api.Worksheet.model_validate(app.get(run_id, "/worksheet").json())
    counts = app.get(run_id, "").json()["counts"]
    assert counts["chains"] == len(dump.chains)


def test_excel_status_shows_override_state_and_falls_back_to_not_credible(
    app: Api, run_id: str
) -> None:
    """Bhargavi, 2026-10-20 (PR 6 answer 1): the Excel Status column is the review state the
    meeting sorts by; the API row keeps `status` and `override_state` apart."""
    findings = [i.target_id for i in app.shortlist(run_id).items if i.target_kind == "finding"]
    accepted, edited, rejected = findings[:3]
    assert app.override(run_id, accepted, "accept", None).status_code == 200
    assert app.override(run_id, edited, "edit", {"consequence": "worse"}).status_code == 200
    assert app.override(run_id, rejected, "reject", None).status_code == 200
    rows = api.Worksheet.model_validate(app.get(run_id, "/worksheet").json()).rows
    sheet = load_workbook(io.BytesIO(app.get(run_id, "/export.xlsx").content)).active
    assert sheet is not None
    table = list(sheet.iter_rows(values_only=True))
    status = [title for title, _ in WORKSHEET_COLUMNS].index("Status")
    shown = {row.row_id: line[status] for row, line in zip(rows, table[1:], strict=True)}
    assert shown[accepted] == "accepted"
    assert shown[edited] == "edited"
    assert shown[rejected] == "rejected"
    untouched = [r for r in rows if r.override_state == "none"]
    dismissed = [r for r in untouched if r.status == "not credible"]
    pending = [r for r in untouched if r.status != "not credible"]
    assert dismissed and pending
    assert {shown[r.row_id] for r in dismissed} == {"not credible — engineer to review"}
    assert {shown[r.row_id] for r in pending} == {"pending review"}
    # the API row is unchanged: the question's status and the override state, apart
    by_id = {r.row_id: r for r in rows}
    assert by_id[accepted].status == "answered" and by_id[accepted].override_state == "accepted"
